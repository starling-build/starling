#!/usr/bin/env python3
"""Generate the fictional, editable Office feature showcases.

Install build/tools/office-examples-requirements.txt in a virtual environment.
Run from anywhere; output defaults to web/office/examples. No app build needed.
"""

import argparse
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree as E
from PIL import Image, ImageDraw
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as DR
from pptx import Presentation
from pptx.util import Inches as PI, Pt as PP
from pptx.dml.color import RGBColor as PC
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.oxml import parse_xml
import xlsxwriter

ROOT = Path(__file__).resolve().parents[2]
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
DGM = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
DSP = "http://schemas.microsoft.com/office/drawing/2008/diagram"
BLUE, INK, GREEN, ROSE = "2449DF", "192B40", "17644F", "AC4328"


def garden_picture():
    """An original site-plan illustration used to demonstrate embedded pictures."""
    im = Image.new("RGB", (1200, 420), "#EDF4EB")
    d = ImageDraw.Draw(im)
    d.rounded_rectangle((20, 20, 1180, 400), radius=25, fill="#D9E8D4")
    d.rectangle((30, 190, 1170, 230), fill="#E8CCA5")
    d.rectangle((565, 25, 635, 395), fill="#E8CCA5")
    for y in [48, 257]:
        for x in [68, 388, 708]:
            d.rounded_rectangle((x, y, x + 240, y + 115), radius=12, fill="#A47552")
            d.rounded_rectangle(
                (x + 9, y + 9, x + 231, y + 106), radius=6, fill="#65503B"
            )
            for u in range(4):
                for v in range(2):
                    cx, cy = x + 40 + u * 52, y + 37 + v * 43
                    d.ellipse((cx - 13, cy - 8, cx + 3, cy + 8), fill="#93BE74")
                    d.ellipse((cx - 3, cy - 8, cx + 13, cy + 8), fill="#5B945E")
    d.ellipse((1030, 55, 1140, 165), fill="#669968")
    d.ellipse((1030, 265, 1140, 375), fill="#669968")
    out = BytesIO()
    im.save(out, format="PNG")
    return out.getvalue()


def word_field(p, code):
    f = OxmlElement("w:fldSimple")
    f.set(qn("w:instr"), code)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    f.append(r)
    p._p.append(f)


def word_link(p, text, url):
    rel = p.part.relate_to(url, DR.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rel)
    r = OxmlElement("w:r")
    pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    pr.append(underline)
    r.append(pr)
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    h.append(r)
    p._p.append(h)


def word_list(doc, text, numbered=False):
    p = doc.add_paragraph(text, "List Number" if numbered else "List Bullet")
    pr = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), "5" if numbered else "1")
    pr.extend([level, num])
    p._p.get_or_add_pPr().append(pr)
    return p


def writer(out, picture):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(0.65)
    sec.left_margin = sec.right_margin = Inches(0.75)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    for style in ["Title", "Heading 1", "Heading 2"]:
        doc.styles[style].font.color.rgb = RGBColor.from_string(BLUE)
    doc.styles["Title"].font.size = Pt(34)
    doc.styles["Heading 1"].font.size = Pt(20)
    doc.styles["Heading 2"].font.size = Pt(14)
    h = sec.header.paragraphs[0]
    h.text = "FIELD NOTES   /   COMMUNITY GARDEN"
    h.style = "Caption"
    footer = sec.footer.paragraphs[0]
    footer.add_run("GARDEN HANDBOOK   •   ")
    word_field(footer, "PAGE")
    doc.add_paragraph("A little room to grow.", "Title")
    doc.add_paragraph("A community garden, built together.", "Subtitle")
    doc.add_picture(BytesIO(picture), width=Inches(7))
    doc.add_paragraph(
        "Six raised beds, a shared path, and room for everyone.", "Caption"
    )
    doc.add_heading("The idea", 1)
    p = doc.add_paragraph("Turn an unused corner into a shared place to ")
    p.add_run("plant, learn, and belong.").bold = True
    p.add_run(" Start with six beds, shared tools, and a weekly open morning.")
    for text in [
        "Grow seasonal food and share what we harvest.",
        "Make first-time gardeners feel at home.",
        "Build a space that gets better with every season.",
    ]:
        word_list(doc, text)
    doc.add_heading("A living plan", 2)
    p = doc.add_paragraph("This handbook is a starting point. ")
    p.add_run("Every neighbor can shape what comes next.").italic = True
    doc.add_paragraph(
        "Try it: change a heading style, resize the picture, or add a bullet. Use the navigation pane to jump between headings.",
        "Quote",
    )
    doc.add_page_break()
    doc.add_heading("From a sketch to a shared space", 1)
    doc.add_paragraph(
        "Four weeks to opening day. Keep the plan practical, visible, and easy to change."
    )
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Shading Accent 1"
    table.autofit = False
    rows = [
        ["Step", "Focus", "Owner", "Timing"],
        ["Listen", "Meet neighbors; gather ideas", "Community team", "Week 1"],
        ["Build", "Prepare beds, paths, and signs", "Build team", "Week 2"],
        ["Plant", "Sow seeds and label each bed", "Garden team", "Week 3"],
        ["Open", "Welcome everyone; share the plan", "All volunteers", "Week 4"],
    ]
    for ri, vals in enumerate(rows):
        cells = table.rows[0].cells if ri == 0 else table.add_row().cells
        for c, t, w in zip(cells, vals, [0.8, 3.1, 1.6, 1.1]):
            c.text = t
            c.width = Inches(w)
    for col, w in zip(table.columns, [0.8, 3.1, 1.6, 1.1]):
        col.width = Inches(w)
    doc.add_heading("Opening morning", 2)
    for text in [
        "Walk the site together and explain the shared tools.",
        "Plant the first seedlings and record what goes where.",
        "Agree on the next watering shift before everyone leaves.",
    ]:
        word_list(doc, text, True)
    doc.add_heading("Everyone has a part", 2)
    p = doc.add_paragraph("Bring gloves, a water bottle, and a neighbor. ")
    r = p.add_run("No experience needed.")
    r.bold = True
    p = doc.add_paragraph("Related tools: ")
    word_link(p, "open the garden budget in Sheets", "https://sheets.starling.build/")
    doc.add_paragraph(
        "Try it: add a schedule row, edit a cell, continue the numbered list, or use Find and Replace to rename a team. Undo restores your last change.",
        "Quote",
    )
    doc.add_page_break()
    doc.add_heading("Care for the season ahead", 1)
    doc.add_heading("Water thoughtfully", 2)
    doc.add_paragraph(
        "Check the soil before watering. A short check each morning is better than a fixed rule that ignores the weather."
    )
    doc.add_heading("Keep useful field notes", 2)
    p = doc.add_paragraph("Write dates, varieties, and what you notice. Mark a ")
    r = p.add_run("small experiment")
    r.italic = True
    p.add_run(" and compare the result next week.")
    p = doc.add_paragraph("A few details can carry meaning: bed area in m")
    p.add_run("2").font.superscript = True
    p.add_run(", water as H")
    p.add_run("2").font.subscript = True
    p.add_run("O, and ")
    r = p.add_run("a revised plan")
    r.font.strike = True
    p.add_run(" a better plan.")
    table = doc.add_table(rows=4, cols=2)
    table.style = "Light Shading Accent 1"
    table.autofit = False
    table.cell(0, 0).merge(table.cell(0, 1)).text = "A shared agreement"
    for r, vals in enumerate(
        [
            (
                "Share the harvest",
                "Leave a portion for neighbors who cannot join the work.",
            ),
            ("Share the tools", "Clean them, return them, and report anything broken."),
            (
                "Share the learning",
                "Keep a short note of what worked and what to try next.",
            ),
        ],
        1,
    ):
        for c, text in zip(table.rows[r].cells, vals):
            c.text = text
    doc.add_paragraph(
        "“A garden grows twice: once in the ground, and once between the people who tend it.”",
        "Quote",
    )
    doc.add_heading("Make this handbook yours", 2)
    doc.add_paragraph(
        "Try the header and footer, page number field, hyperlink, table formatting, text color, and alignment. Save a copy as DOCX, then reopen it to continue."
    )
    doc.core_properties.title = "A little room to grow — community garden handbook"
    doc.core_properties.author = "Starling Office"
    doc.save(out / "garden-brief.docx")


def slide_text(s, x, y, w, h, text, size=22, color=INK, bold=False):
    shape = s.shapes.add_textbox(PI(x), PI(y), PI(w), PI(h))
    tf = shape.text_frame
    tf.word_wrap = True
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = "Calibri"
        p.font.size = PP(size)
        p.font.color.rgb = PC.from_string(color)
        p.font.bold = bold
    return shape


def slide_shape(s, x, y, w, h, color, kind=MSO_SHAPE.RECTANGLE, text=None):
    sh = s.shapes.add_shape(kind, PI(x), PI(y), PI(w), PI(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = PC.from_string(color)
    sh.line.fill.background()
    if text:
        sh.text = text
        for p in sh.text_frame.paragraphs:
            p.font.size = PP(22)
            p.font.color.rgb = PC.from_string("FFFFFF")
    return sh


def entrance(slide, shape_id):
    # A standard PowerPoint click-to-wipe entrance, understood by AnimationXML.
    xml = f"""<p:timing xmlns:p="{P}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst><p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst><p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst><p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst><p:par><p:cTn id="5" presetID="22" presetClass="entr" presetSubtype="8" fill="hold" nodeType="clickEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst><p:set><p:cBhvr><p:cTn id="6" dur="1" fill="hold"/><p:tgtEl><p:spTgt spid="{shape_id}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set><p:animEffect transition="in" filter="wipe(right)"><p:cBhvr><p:cTn id="7" dur="700"/><p:tgtEl><p:spTgt spid="{shape_id}"/></p:tgtEl></p:cBhvr></p:animEffect></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq></p:childTnLst></p:cTn></p:par></p:tnLst><p:bldLst><p:bldP spid="{shape_id}" grpId="0" animBg="1"/></p:bldLst></p:timing>"""
    slide._element.append(parse_xml(xml))


def slides(out, picture):
    prs = Presentation()
    prs.slide_width = PI(13.333)
    prs.slide_height = PI(7.5)

    def page(title, number, tip):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = PC.from_string("FFF6EF")
        slide_shape(s, 0, 0, 0.16, 7.5, "DC623D")
        slide_text(
            s,
            0.65,
            0.35,
            12,
            0.5,
            f"FIELD NOTES / COMMUNITY GARDEN   /   {number:02}",
            12,
            ROSE,
            True,
        )
        slide_text(s, 0.65, 1.05, 12, 1, title, 32, bold=True)
        slide_text(s, 0.65, 6.85, 12, 0.4, tip, 12, ROSE)
        s.notes_slide.notes_text_frame.text = (
            tip
            + " These are fictional planning figures. Invite the audience to change the plan."
        )
        return s

    def chart(s, kind, x, y, w, h, title, series=None):
        data = CategoryChartData()
        data.categories = ["Prepare", "Plant", "Grow"]
        for name, values in series or [("Volunteer hours", [18, 32, 48])]:
            data.add_series(name, values)
        c = s.shapes.add_chart(kind, PI(x), PI(y), PI(w), PI(h), data).chart
        c.has_title = True
        c.chart_title.text_frame.text = title
        c.has_legend = False
        for i, se in enumerate(c.series):
            se.format.fill.solid()
            se.format.fill.fore_color.rgb = PC.from_string(["DC623D", "17644F"][i % 2])
        return c

    s = page(
        "Good things\ngrow together.",
        1,
        "Try it: edit the chart data. Start Slide Show to present; open Presenter View for notes.",
    )
    # Give the cover title more height and a narrower measure.
    title = s.shapes[2]
    title.width = PI(6)
    title.height = PI(2.4)
    for p in title.text_frame.paragraphs:
        p.font.size = PP(42)
    slide_text(
        s,
        0.65,
        3.8,
        5.8,
        1.2,
        "A shared space. A simple plan.\nA whole season of possibility.",
        22,
    )
    for x, n, label in [
        (0.65, "06", "raised beds"),
        (2.65, "04", "weeks to open"),
        (4.65, "01", "community"),
    ]:
        slide_text(s, x, 5.4, 1.8, 0.6, n, 30, ROSE, True)
        slide_text(s, x, 6, 2, 0.4, label, 14)
    chart(
        s,
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        7,
        1.7,
        5.7,
        4.9,
        "More hands. More possibilities.",
    )
    s = page(
        "Make room for everyone.",
        2,
        "Try it: move or resize the picture, edit the caption, and add a link to the text.",
    )
    s.shapes.add_picture(BytesIO(picture), PI(0.7), PI(2.4), width=PI(8.1))
    slide_text(
        s,
        9.25,
        2.5,
        3.1,
        3.5,
        "Six raised beds.\nShared tools.\nOpen mornings.\nSpace to learn.",
        24,
    )
    slide_text(s, 0.7, 5.7, 11, 0.5, "A simple site plan, embedded as a picture.", 18)
    s = page(
        "Four weeks to a fresh start.",
        3,
        "Try it: edit and align the shapes. In Slide Show, click to reveal the final step with a wipe.",
    )
    for i, (label, color) in enumerate(
        [
            ("Listen", "2449DF"),
            ("Build", "17644F"),
            ("Plant", "AC4328"),
            ("Grow", "7150A0"),
        ]
    ):
        sh = slide_shape(
            s, 0.7 + i * 3.15, 2.7, 2.7, 1.7, color, MSO_SHAPE.CHEVRON, label
        )
    entrance(s, sh.shape_id)
    slide_text(
        s, 0.8, 5.1, 11.7, 0.8, "One shared plan. A little progress every week.", 24
    )
    s._element.insert(
        len(s._element) - 1,
        parse_xml(f'<p:transition xmlns:p="{P}" spd="med"><p:fade/></p:transition>'),
    )
    s = page(
        "A plan people can follow.",
        4,
        "Try it: edit table cells, change fills, and adjust row heights and column widths.",
    )
    table = s.shapes.add_table(5, 4, PI(0.7), PI(2.35), PI(11.9), PI(3.8)).table
    for ri, row in enumerate(
        [
            ["Week", "Work", "People", "Result"],
            ["1", "Listen and sketch", "Neighbors", "A shared plan"],
            ["2", "Build beds and paths", "Build team", "A welcoming space"],
            ["3", "Plant and label", "Garden team", "The first seedlings"],
            ["4", "Open the gates", "Everyone", "A new habit"],
        ]
    ):
        for ci, value in enumerate(row):
            cell = table.cell(ri, ci)
            cell.text = value
            cell.fill.solid()
            cell.fill.fore_color.rgb = PC.from_string(
                "AC4328" if ri == 0 else ("F4DDCE" if ri % 2 else "FFF6EF")
            )
            for p in cell.text_frame.paragraphs:
                p.font.size = PP(19)
                p.font.color.rgb = PC.from_string("FFFFFF" if ri == 0 else INK)
    table.columns[0].width = PI(1.1)
    table.columns[1].width = PI(4)
    table.columns[2].width = PI(3.2)
    table.columns[3].width = PI(3.6)
    s = page(
        "Learn. Share. Grow.",
        5,
        "Try it: select the SmartArt diagram, move or resize it, then save and reopen the deck.",
    )
    # The diagram's editable source data and cached DrawingML are added below.
    frame = s.shapes.add_table(1, 1, PI(0.7), PI(2.45), PI(11.9), PI(3.5))._element
    gd = frame.find(f".//{{{A}}}graphicData")
    gd.clear()
    gd.set("uri", DGM)
    gd.append(
        parse_xml(
            f'<dgm:relIds xmlns:dgm="{DGM}" xmlns:r="{R}" r:dm="rIdGardenData" r:lo="rIdGardenLayout" r:qs="rIdGardenStyle" r:cs="rIdGardenColors"/>'
        )
    )
    s = page(
        "A season of small gains.",
        6,
        "Try it: change the line and area chart values to explore a different season.",
    )
    chart(s, XL_CHART_TYPE.LINE, 0.6, 2.15, 6, 4.3, "Weekly participation")
    chart(s, XL_CHART_TYPE.AREA, 6.8, 2.15, 5.9, 4.3, "Cumulative harvest")
    s = page(
        "See the balance.",
        7,
        "Try it: compare a bar chart with a pie chart. Change the chart colors and labels.",
    )
    chart(s, XL_CHART_TYPE.BAR_CLUSTERED, 0.6, 2.15, 6, 4.3, "Hours by activity")
    chart(s, XL_CHART_TYPE.PIE, 6.8, 2.15, 5.9, 4.3, "Share of the work")
    s = page(
        "Look for a pattern.",
        8,
        "Try it: edit a scatter plot point or a doughnut segment. All charts contain real data.",
    )
    data = XyChartData()
    series = data.add_series("Beds")
    for x, y in [(2, 5), (4, 12), (6, 18), (8, 25), (10, 29)]:
        series.add_data_point(x, y)
    c = s.shapes.add_chart(
        XL_CHART_TYPE.XY_SCATTER, PI(0.6), PI(2.15), PI(6), PI(4.3), data
    ).chart
    c.has_title = True
    c.chart_title.text_frame.text = "Watering time and harvest"
    c.has_legend = False
    chart(s, XL_CHART_TYPE.DOUGHNUT, 6.8, 2.15, 5.9, 4.3, "Where the time goes")
    s = page(
        "Let’s grow something good.",
        9,
        "Try it: change the WordArt, rotate the star, or add a new slide using a different layout.",
    )
    sh = slide_text(s, 0.7, 2.4, 11.5, 1.8, "GROW TOGETHER", 48, GREEN, True)
    sh.text_frame._txBody.bodyPr.append(
        parse_xml(
            f'<a:prstTxWarp xmlns:a="{A}" prst="textArchUp"><a:avLst/></a:prstTxWarp>'
        )
    )
    star = slide_shape(s, 10.6, 4.4, 1.35, 1.35, "DC623D", MSO_SHAPE.STAR_5_POINT)
    star.rotation = 15
    sh = slide_text(
        s, 0.8, 4.8, 9, 1.1, "Bring an idea. Bring a neighbor. Start here.", 26
    )
    sh.text_frame.paragraphs[0].runs[
        0
    ].hyperlink.address = "https://writer.starling.build/"
    prs.save(out / "garden-presentation.pptx")
    add_diagram(out / "garden-presentation.pptx")


def add_diagram(path):
    """A three-node SmartArt process, including data and a cached vector drawing."""
    with ZipFile(path) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    relns = "http://schemas.openxmlformats.org/package/2006/relationships"
    rel = E.fromstring(parts["ppt/slides/_rels/slide5.xml.rels"])
    for suffix, file, kind in [
        ("Data", "data1.xml", "diagramData"),
        ("Layout", "layout1.xml", "diagramLayout"),
        ("Style", "quickStyle1.xml", "diagramQuickStyle"),
        ("Colors", "colors1.xml", "diagramColors"),
        ("Drawing", "drawing1.xml", "diagramDrawing"),
    ]:
        kind_uri = (
            (R + "/" + kind)
            if kind != "diagramDrawing"
            else "http://schemas.microsoft.com/office/2007/relationships/diagramDrawing"
        )
        E.SubElement(
            rel,
            "{" + relns + "}Relationship",
            Id="rIdGarden" + suffix,
            Type=kind_uri,
            Target="../diagrams/" + file,
        )
    parts["ppt/slides/_rels/slide5.xml.rels"] = E.tostring(rel)
    parts["ppt/diagrams/_rels/data1.xml.rels"] = (
        f'<Relationships xmlns="{relns}"><Relationship Id="rIdGardenDrawing" Type="http://schemas.microsoft.com/office/2007/relationships/diagramDrawing" Target="drawing1.xml"/></Relationships>'.encode()
    )
    nodes = ""
    connections = ""
    for i, label in enumerate(["Learn", "Share", "Grow"], 1):
        nodes += f'<dgm:pt modelId="{i}"><dgm:prSet/><dgm:spPr/><dgm:t><a:bodyPr/><a:lstStyle/><a:p><a:r><a:t>{label}</a:t></a:r></a:p></dgm:t></dgm:pt>'
        connections += f'<dgm:cxn modelId="{i+10}" type="parOf" srcId="0" destId="{i}" srcOrd="{i-1}" destOrd="0"/>'
    parts["ppt/diagrams/data1.xml"] = (
        f'<dgm:dataModel xmlns:dgm="{DGM}" xmlns:a="{A}" xmlns:dsp="{DSP}"><dgm:ptLst><dgm:pt modelId="0" type="doc"><dgm:prSet/></dgm:pt>{nodes}</dgm:ptLst><dgm:cxnLst>{connections}</dgm:cxnLst><dgm:bg/><dgm:whole/><dgm:extLst><a:ext uri="http://schemas.microsoft.com/office/drawing/2008/diagram"><dsp:dataModelExt relId="rIdGardenDrawing" minVer="http://schemas.openxmlformats.org/drawingml/2006/diagram"/></a:ext></dgm:extLst></dgm:dataModel>'.encode()
    )
    # A real diagram drawing: individual vector nodes, not a flattened picture.
    shapes = ""
    for i, (label, color) in enumerate(
        zip(["Learn", "Share", "Grow"], [BLUE, GREEN, ROSE])
    ):
        shapes += f'<dsp:sp modelId="{i+1}"><dsp:nvSpPr><dsp:cNvPr id="{i+2}" name="{label}"/><dsp:cNvSpPr/><dsp:nvPr/></dsp:nvSpPr><dsp:spPr><a:xfrm><a:off x="{i*3600000}" y="0"/><a:ext cx="3200000" cy="2000000"/></a:xfrm><a:prstGeom prst="chevron"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val="{color}"/></a:solidFill><a:ln><a:noFill/></a:ln></dsp:spPr><dsp:txBody><a:bodyPr anchor="ctr"/><a:lstStyle/><a:p><a:pPr algn="ctr"/><a:r><a:rPr sz="3000"><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill></a:rPr><a:t>{label}</a:t></a:r></a:p></dsp:txBody></dsp:sp>'
    parts["ppt/diagrams/drawing1.xml"] = (
        f'<dsp:drawing xmlns:dsp="{DSP}" xmlns:a="{A}"><dsp:spTree><dsp:nvGrpSpPr><dsp:cNvPr id="1" name=""/><dsp:cNvGrpSpPr/><dsp:nvPr/></dsp:nvGrpSpPr><dsp:grpSpPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="10400000" cy="2000000"/><a:chOff x="0" y="0"/><a:chExt cx="10400000" cy="2000000"/></a:xfrm></dsp:grpSpPr>{shapes}</dsp:spTree></dsp:drawing>'.encode()
    )
    # Standard part roots; cached drawing defines the authored presentation.
    parts["ppt/diagrams/layout1.xml"] = (
        f'<dgm:layoutDef xmlns:dgm="{DGM}" uniqueId="urn:starling:garden:process" minVer="12.0"><dgm:title val="Garden process"/><dgm:desc val="Learn, share, grow"/><dgm:catLst/><dgm:layoutNode name="process"><dgm:alg type="lin"/><dgm:forEach name="nodes" axis="ch" ptType="node"><dgm:layoutNode name="node"><dgm:alg type="tx"/><dgm:shape type="chevron"/></dgm:layoutNode></dgm:forEach></dgm:layoutNode></dgm:layoutDef>'.encode()
    )
    parts["ppt/diagrams/quickStyle1.xml"] = (
        f'<dgm:styleDef xmlns:dgm="{DGM}" uniqueId="urn:starling:garden:style" minVer="12.0"><dgm:title val="Garden"/><dgm:desc val=""/><dgm:catLst/></dgm:styleDef>'.encode()
    )
    parts["ppt/diagrams/colors1.xml"] = (
        f'<dgm:colorsDef xmlns:dgm="{DGM}" uniqueId="urn:starling:garden:colors" minVer="12.0"><dgm:title val="Garden"/><dgm:desc val=""/><dgm:catLst/></dgm:colorsDef>'.encode()
    )
    ct = E.fromstring(parts["[Content_Types].xml"])
    ctns = ct.nsmap[None]
    for file, kind in [
        ("data1.xml", "Data"),
        ("layout1.xml", "Layout"),
        ("quickStyle1.xml", "Style"),
        ("colors1.xml", "Colors"),
    ]:
        E.SubElement(
            ct,
            "{" + ctns + "}Override",
            PartName="/ppt/diagrams/" + file,
            ContentType="application/vnd.openxmlformats-officedocument.drawingml.diagram"
            + kind
            + "+xml",
        )
    E.SubElement(
        ct,
        "{" + ctns + "}Override",
        PartName="/ppt/diagrams/drawing1.xml",
        ContentType="application/vnd.ms-office.drawingml.diagramDrawing+xml",
    )
    parts["[Content_Types].xml"] = E.tostring(ct)
    with ZipFile(path, "w", ZIP_DEFLATED) as z:
        for n, data in parts.items():
            z.writestr(n, data)


def sheets(out):
    wb = xlsxwriter.Workbook(out / "garden-budget.xlsx", {"strings_to_formulas": False})
    wb.set_properties(
        {"title": "A community garden, in numbers", "author": "Starling Office"}
    )

    def fmt(**kw):
        return wb.add_format(
            {"font_name": "Calibri", "font_size": 12, "font_color": "#243E34", **kw}
        )

    title = fmt(font_size=22, bold=True, font_color="#135F49")
    subtitle = fmt(font_size=11, font_color="#517267")
    head = fmt(bold=True, bg_color="#17644F", font_color="white")
    body = fmt()
    money = fmt(num_format="$#,##0")
    percent = fmt(num_format="0%")
    total = fmt(bold=True, bg_color="#DDEEE4", num_format="$#,##0")
    date = fmt(num_format="mmm d")
    inputfmt = fmt(bg_color="#FFF0CC", locked=False, num_format="0")

    def page(name, heading, sub, cols="A:F"):
        ws = wb.add_worksheet(name)
        ws.hide_gridlines(2)
        ws.set_default_row(25)
        ws.set_column(cols, 18)
        ws.merge_range("A1:F1", heading, title)
        ws.set_row(0, 40)
        ws.merge_range("A2:F2", sub, subtitle)
        ws.set_row(1, 30)
        ws.set_landscape()
        ws.set_paper(9)
        ws.fit_to_pages(1, 1)
        ws.set_margins(0.3, 0.3, 0.5, 0.5)
        ws.set_header("&LField Notes / Community Garden")
        ws.set_footer("&RPage &P of &N")
        return ws

    ws = page(
        "Garden budget",
        "A community garden, in numbers.",
        "Edit an actual amount. Totals, status colors, and the chart follow.",
    )
    ws.set_column("A:A", 26)
    rows = [
        ["Raised beds", 2400, 2150],
        ["Soil & compost", 1200, 1080],
        ["Seeds & plants", 850, 920],
        ["Shared tools", 1100, 960],
        ["Paths & signs", 1450, 1320],
        ["Opening day", 500, 390],
    ]
    data = []
    for i, (name, b, a) in enumerate(rows, 5):
        data.append(
            [name, b, a, f"=B{i}-C{i}", f'=IF(D{i}<0,"Over budget","On track")']
        )
    ws.add_table(
        "A4:E10",
        {
            "name": "GardenBudget",
            "style": "Table Style Medium 4",
            "data": data,
            "columns": [
                {"header": "Category"},
                {"header": "Budget", "format": money},
                {"header": "Actual", "format": money},
                {"header": "Remaining", "format": money},
                {"header": "Status"},
            ],
        },
    )
    for i, (_, b, a) in enumerate(rows, 4):
        ws.write_formula(i, 3, f"=B{i+1}-C{i+1}", money, b - a)
        ws.write_formula(
            i,
            4,
            f'=IF(D{i+1}<0,"Over budget","On track")',
            body,
            "Over budget" if b < a else "On track",
        )
    ws.write("A12", "Total", head)
    for col, formula, result in [
        ("B", "=SUM(B5:B10)", 7500),
        ("C", "=SUM(C5:C10)", 6820),
        ("D", "=SUM(D5:D10)", 680),
    ]:
        ws.write_formula(col + "12", formula, total, result)
    ws.conditional_format(
        "D5:D10",
        {
            "type": "cell",
            "criteria": "<",
            "value": 0,
            "format": fmt(bg_color="#F8D8D6", font_color="#A32921"),
        },
    )
    ws.conditional_format("D5:D10", {"type": "data_bar", "bar_color": "#7DBBA0"})
    ws.data_validation(
        "C5:C10",
        {
            "validate": "decimal",
            "criteria": ">=",
            "value": 0,
            "input_title": "Actual spending",
            "input_message": "Enter a non-negative amount.",
            "error_message": "Use a number greater than or equal to zero.",
        },
    )
    ws.freeze_panes(4, 1)
    ws.merge_range(
        "A15:E15",
        "Try the table filters, sort by budget, or change an actual amount.",
        subtitle,
    )
    c = wb.add_chart({"type": "column"})
    for col, color in [("B", "#17644F"), ("C", "#7DBBA0")]:
        c.add_series(
            {
                "name": f"='Garden budget'!${col}$4",
                "categories": "='Garden budget'!$A$5:$A$10",
                "values": f"='Garden budget'!${col}$5:${col}$10",
                "fill": {"color": color},
            }
        )
    c.set_title({"name": "Budget and actual"})
    c.set_size({"width": 580, "height": 360})
    ws.insert_chart("G4", c)
    ws.print_area("A1:O18")
    ws = page(
        "Season",
        "Watch the season take shape.",
        "Dates, percentages, grouped rows, and charts. Expand the outline to inspect a month.",
    )
    ws.write_row(
        "A4", ["Month", "Visits", "Harvest kg", "Target kg", "Progress", "Label"], head
    )
    import datetime

    for i, (visits, harvest, target) in enumerate(
        [
            (24, 12, 20),
            (32, 22, 25),
            (44, 35, 30),
            (50, 42, 40),
            (55, 48, 45),
            (60, 54, 50),
        ],
        4,
    ):
        ws.write_datetime(i, 0, datetime.datetime(2026, i + 1, 1), date)
        ws.write_formula(
            i,
            5,
            f'=TEXT(A{i+1},"mmm")',
            body,
            ["May", "Jun", "Jul", "Aug", "Sep", "Oct"][i - 4],
        )
        ws.write_row(i, 1, [visits, harvest, target], body)
        ws.write_formula(i, 4, f"=C{i+1}/D{i+1}", percent, harvest / target)
    ws.set_row(5, 25, None, {"level": 1})
    ws.set_row(6, 25, None, {"level": 1})
    ws.conditional_format("E5:E10", {"type": "3_color_scale"})
    ws.freeze_panes(4, 1)
    c = wb.add_chart({"type": "line"})
    c.add_series(
        {
            "name": "Harvest",
            "categories": "='Season'!$F$5:$F$10",
            "values": "='Season'!$C$5:$C$10",
            "line": {"color": "#17644F"},
        }
    )
    c.set_title({"name": "Harvest through the season"})
    c.set_size({"width": 560, "height": 320})
    ws.insert_chart("G4", c)
    ws.print_area("A1:O18")
    ws = page(
        "Volunteers",
        "Many hands, shared work.",
        "Filter a team, sort by hours, or choose a different role from the dropdown.",
    )
    ws.set_column("A:A", 23)
    volunteers = [
        ["Alex", "Build", 6, "Morning"],
        ["Sam", "Garden", 4, "Afternoon"],
        ["Jordan", "Welcome", 3, "Morning"],
        ["Robin", "Garden", 5, "Morning"],
        ["Casey", "Build", 4, "Afternoon"],
        ["Morgan", "Welcome", 2, "Afternoon"],
    ]
    ws.add_table(
        "A4:D10",
        {
            "name": "Volunteers",
            "style": "Table Style Medium 4",
            "data": volunteers,
            "columns": [{"header": x} for x in ["Name", "Team", "Hours", "Shift"]],
        },
    )
    ws.data_validation(
        "B5:B10", {"validate": "list", "source": ["Build", "Garden", "Welcome"]}
    )
    ws.data_validation(
        "C5:C10",
        {"validate": "integer", "criteria": "between", "minimum": 0, "maximum": 12},
    )
    ws.freeze_panes(4, 1)
    ws.write_url("A13", "internal:'Explore'!A1", string="Explore the formulas →")
    c = wb.add_chart({"type": "pie"})
    c.add_series(
        {
            "name": "Volunteer hours",
            "categories": "='Volunteers'!$A$5:$A$10",
            "values": "='Volunteers'!$C$5:$C$10",
        }
    )
    c.set_title({"name": "Sharing the work"})
    c.set_size({"width": 480, "height": 320})
    ws.insert_chart("G4", c)
    ws.print_area("A1:N18")
    ws = page(
        "Explore",
        "Ask a better question.",
        "Live formulas: edit a source sheet, then see these answers change.",
    )
    ws.set_column("A:A", 32)
    ws.set_column("B:B", 22)
    ws.set_column("C:F", 19)
    ws.write_row("A4", ["Question", "Answer", "Formula to explore"], head)
    formulas = [
        ("Total volunteer hours", "=SUM(Volunteers!C5:C10)", 24),
        (
            "Garden team hours",
            '=SUMIF(Volunteers!B5:B10,"Garden",Volunteers!C5:C10)',
            9,
        ),
        ("Morning volunteers", '=COUNTIF(Volunteers!D5:D10,"Morning")', 3),
        ("Average shift", "=AVERAGE(Volunteers!C5:C10)", 4),
        ("Largest shift", "=MAX(Volunteers!C5:C10)", 6),
        ("Rounded mean", "=ROUND(AVERAGE(Volunteers!C5:C10),1)", 4),
        (
            "Alex’s team",
            '=XLOOKUP("Alex",Volunteers!A5:A10,Volunteers!B5:B10)',
            "Build",
        ),
        ("A safe fallback", '=IFERROR(1/0,"Try another value")', "Try another value"),
        (
            "Name an intermediate value",
            "=LET(hours,SUM(Volunteers!C5:C10),hours*2)",
            48,
        ),
        ("A reusable calculation", "=DoubleHours(12)", 24),
    ]
    wb.define_name("DoubleHours", "=LAMBDA(hours,hours*2)")
    wb.define_name("TotalBudget", "='Garden budget'!$B$12")
    for r, (question, formula, value) in enumerate(formulas, 4):
        ws.write(r, 0, question, body)
        ws.write_formula(r, 1, formula, body, value)
        ws.merge_range(r, 2, r, 5, formula, subtitle)
    # Dynamic arrays have empty spill ranges so visitors can change their sources.
    ws.write("A17", "Unique teams", head)
    ws.write_dynamic_array_formula(
        "A18:A20", "=UNIQUE(Volunteers!B5:B10)", body, "Build"
    )
    ws.write("C17", "A four-week sequence", head)
    ws.write_dynamic_array_formula("C18:C21", "=SEQUENCE(4)", body, 1)
    ws.write("E17", "Garden volunteers", head)
    ws.write_dynamic_array_formula(
        "E18:E19", '=FILTER(Volunteers!A5:A10,Volunteers!B5:B10="Garden")', body, "Sam"
    )
    ws.write("A24", "Names in order", head)
    ws.write_dynamic_array_formula("A25:A30", "=SORT(Volunteers!A5:A10)", body, "Alex")
    ws.freeze_panes(4, 1)
    ws.print_area("A1:F31")
    ws = page(
        "Protected plan",
        "Change the inputs. Keep the plan.",
        "Yellow cells are editable. Formula cells are protected without a password.",
    )
    ws.set_column("A:A", 30)
    ws.write_row("A4", ["Input", "Value", "Meaning"], head)
    for r, label, value, meaning in [
        (4, "Beds", 6, "Raised beds to prepare"),
        (5, "Plants per bed", 12, "Your planting density"),
        (6, "Cost per plant", 3, "Illustrative unit cost"),
    ]:
        ws.write(r, 0, label, body)
        ws.write(r, 1, value, inputfmt)
        ws.merge_range(r, 2, r, 5, meaning, body)
    ws.write("A9", "Plants needed", head)
    ws.write_formula("B9", "=B5*B6", body, 72)
    ws.write("A10", "Planting cost", head)
    ws.write_formula("B10", "=B9*B7", money, 216)
    ws.write("A11", "Share of total budget", head)
    ws.write_formula("B11", "=B10/TotalBudget", percent, 216 / 7500)
    ws.data_validation("B5:B7", {"validate": "integer", "criteria": ">", "value": 0})
    ws.merge_range(
        "A14:F14",
        "Try editing a yellow cell, then a protected total. You can unprotect this sheet.",
        subtitle,
    )
    ws.protect("", {"select_locked_cells": True, "select_unlocked_cells": True})
    ws.print_area("A1:F16")
    wb.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=ROOT / "web/office/examples")
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    picture = garden_picture()
    writer(args.output, picture)
    slides(args.output, picture)
    sheets(args.output)
    for path in sorted(args.output.glob("garden-*")):
        print(path)


if __name__ == "__main__":
    main()
